"""Compare a preview Excel artifact with the shadow Feishu baseline.

This grader is the stronger sibling of ``preview_file_exists``. It still starts
from the public case ``outputs/`` directory, but then reads the generated
preview Excel, finds the matching result table from the Feishu project catalog,
fetches the same-period baseline records, and compares the actual row data.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import json
import os
import re
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook
import requests

from evals.graders.preview_file_exists import _find_preview_file, _outputs_relative_path

PERIOD_FIELD_NAME = "期间"
PREVIEW_RESULTS_SHEET = "results"
PREVIEW_METADATA_SHEET = "__meta__"
TEXTUAL_FIELD_TYPES = {1}
NUMERIC_FIELD_TYPES = {2}
MAX_SAMPLE_DIFFERENCES = 5
MAX_SAMPLE_FIELD_DIFFERENCES = 5


@dataclass(frozen=True)
class PreviewArtifactData:
    """Preview Excel content and metadata used by the baseline grader.

    Args:
        records: Rows from the ``results`` sheet, represented as plain dicts.
        metadata: Key/value pairs from the hidden ``__meta__`` sheet.
        headers: Ordered result-sheet headers.

    Returns:
        Immutable preview artifact data.
    """

    records: list[dict[str, Any]]
    metadata: dict[str, Any]
    headers: list[str]


@dataclass(frozen=True)
class DifferenceDetail:
    """One field-level difference between preview and baseline.

    Args:
        field: Field name that differs.
        preview_value: Normalized value from the preview artifact.
        baseline_value: Normalized value from the Feishu baseline.

    Returns:
        Immutable difference detail.
    """

    field: str
    preview_value: Any
    baseline_value: Any


@dataclass(frozen=True)
class DifferenceEntry:
    """One record-level comparison issue.

    Args:
        key: Business key built from ``期间`` plus preview ``group_fields``.
        issue_type: Difference category, such as ``value_mismatch``.
        differences: Field-level details for this record.

    Returns:
        Immutable difference entry.
    """

    key: tuple[Any, ...]
    issue_type: str
    differences: list[DifferenceDetail]


@dataclass(frozen=True)
class BaselineDiffReport:
    """Complete preview/baseline comparison result.

    Args:
        matched: Whether the preview fully matches the baseline.
        preview_count: Number of preview records compared.
        baseline_count: Number of baseline records compared.
        key_fields: Business key fields used for row alignment.
        compared_fields: Fields included in the field-level comparison.
        differences: All detected differences.

    Returns:
        Immutable diff report.
    """

    matched: bool
    preview_count: int
    baseline_count: int
    key_fields: list[str]
    compared_fields: list[str]
    differences: list[DifferenceEntry]


class BaselineSource(Protocol):
    """Read-only source for project table metadata and baseline records."""

    def table_id_for_recog_id(self, recog_id: str) -> str | None:
        """Return the Feishu result table id for one recognition project.

        Args:
            recog_id: Recognition project id from the preview metadata.

        Returns:
            Result-table id, or ``None`` when the project is not registered.
        """

    def fetch_period_records(
        self,
        *,
        table_id: str,
        period: str,
        field_names: list[str],
    ) -> list[dict[str, Any]]:
        """Fetch same-period baseline records from the result table.

        Args:
            table_id: Feishu Bitable table id.
            period: Accounting period in ``YYYYMM`` format.
            field_names: Field names to fetch. The grader intentionally asks
                only for preview fields, so downstream formula-only fields do
                not affect the rollup check.

        Returns:
            Flattened Feishu records.
        """


class GraderConfigurationError(RuntimeError):
    """Raised when the baseline grader lacks required environment settings."""


class FeishuBaselineSource:
    """Read baseline data from Feishu Bitable in a strictly read-only manner."""

    def __init__(
        self,
        *,
        app_id: str,
        app_secret: str,
        config_bitable_app_token: str,
        result_bitable_app_token: str,
        project_catalog_table_id: str,
        timeout_seconds: int = 30,
        trust_env: bool = False,
    ) -> None:
        """Initialize the Feishu baseline source.

        Args:
            app_id: Feishu application id.
            app_secret: Feishu application secret.
            config_bitable_app_token: Bitable app token that stores project
                catalog configuration.
            result_bitable_app_token: Bitable app token that stores shadow
                result tables.
            project_catalog_table_id: Table id for the project catalog.
            timeout_seconds: HTTP timeout in seconds.
            trust_env: Whether requests should inherit proxy settings from the
                host environment.

        Returns:
            None.
        """

        self._config_bitable_app_token = config_bitable_app_token
        self._result_bitable_app_token = result_bitable_app_token
        self._project_catalog_table_id = project_catalog_table_id
        self._timeout_seconds = timeout_seconds
        self._session = requests.Session()
        self._session.trust_env = trust_env
        self._token_provider = _TenantTokenProvider(
            app_id=app_id,
            app_secret=app_secret,
            timeout_seconds=timeout_seconds,
            trust_env=trust_env,
        )

    @classmethod
    def from_env(cls) -> "FeishuBaselineSource":
        """Build a Feishu baseline source from process environment variables.

        Args:
            None.

        Returns:
            Configured FeishuBaselineSource.

        Raises:
            GraderConfigurationError: If required settings are missing.
        """

        required = {
            "RERE_FEISHU_APP_ID": os.environ.get("RERE_FEISHU_APP_ID"),
            "RERE_FEISHU_APP_SECRET": os.environ.get("RERE_FEISHU_APP_SECRET"),
            "RERE_CONFIG_BITABLE_APP_TOKEN": os.environ.get("RERE_CONFIG_BITABLE_APP_TOKEN"),
            "RERE_RESULT_BITABLE_APP_TOKEN": os.environ.get("RERE_RESULT_BITABLE_APP_TOKEN"),
            "RERE_PROJECT_CATALOG_TABLE_ID": os.environ.get("RERE_PROJECT_CATALOG_TABLE_ID"),
        }
        missing = [key for key, value in required.items() if not value]
        if missing:
            raise GraderConfigurationError("missing environment variables: " + ", ".join(missing))

        timeout_seconds = int(os.environ.get("RERE_REQUEST_TIMEOUT_SECONDS", "30"))
        trust_env = os.environ.get("RERE_TRUST_ENV_PROXIES", "false").strip().lower() in {"1", "true", "yes", "y"}
        return cls(
            app_id=required["RERE_FEISHU_APP_ID"] or "",
            app_secret=required["RERE_FEISHU_APP_SECRET"] or "",
            config_bitable_app_token=required["RERE_CONFIG_BITABLE_APP_TOKEN"] or "",
            result_bitable_app_token=required["RERE_RESULT_BITABLE_APP_TOKEN"] or "",
            project_catalog_table_id=required["RERE_PROJECT_CATALOG_TABLE_ID"] or "",
            timeout_seconds=timeout_seconds,
            trust_env=trust_env,
        )

    def table_id_for_recog_id(self, recog_id: str) -> str | None:
        """Return the result table id registered for a recognition project.

        Args:
            recog_id: Recognition project id.

        Returns:
            Result-table id, or ``None`` when no matching project exists.
        """

        records = self._paginate(
            method="GET",
            path=f"/open-apis/bitable/v1/apps/{self._config_bitable_app_token}/tables/{self._project_catalog_table_id}/records",
            list_key="items",
        )
        for record in records:
            fields = record.get("fields", {})
            current_recog_id = _extract_feishu_text(fields.get("recog_id"))
            if current_recog_id == recog_id:
                return _extract_feishu_text(fields.get("bitable_table_id"))
        return None

    def fetch_period_records(
        self,
        *,
        table_id: str,
        period: str,
        field_names: list[str],
    ) -> list[dict[str, Any]]:
        """Fetch baseline records for one period from one result table.

        Args:
            table_id: Feishu result table id.
            period: Accounting period in ``YYYYMM`` format.
            field_names: Field names to return.

        Returns:
            Flattened Feishu records.
        """

        normalized_field_names = _ensure_period_field(field_names)
        period_field_type = self._lookup_period_field_type(table_id=table_id)
        filter_value = _build_period_filter_value(period, period_field_type)
        records = self._paginate(
            method="POST",
            path=f"/open-apis/bitable/v1/apps/{self._result_bitable_app_token}/tables/{table_id}/records/search",
            list_key="items",
            body={
                "field_names": normalized_field_names,
                "filter": {
                    "conjunction": "and",
                    "conditions": [
                        {
                            "field_name": PERIOD_FIELD_NAME,
                            "operator": "is",
                            "value": [filter_value],
                        }
                    ],
                },
            },
        )
        return [flatten_feishu_fields(record.get("fields", {})) for record in records]

    def _lookup_period_field_type(self, *, table_id: str) -> int | None:
        """Read the Feishu field type for ``期间``.

        Args:
            table_id: Feishu result table id.

        Returns:
            Feishu field type id, or ``None`` when unavailable.
        """

        fields = self._paginate(
            method="GET",
            path=f"/open-apis/bitable/v1/apps/{self._result_bitable_app_token}/tables/{table_id}/fields",
            list_key="items",
        )
        for field_item in fields:
            if field_item.get("field_name") == PERIOD_FIELD_NAME:
                field_type = field_item.get("type")
                return field_type if isinstance(field_type, int) else None
        return None

    def _paginate(
        self,
        *,
        method: str,
        path: str,
        list_key: str,
        body: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        """Read all pages from a Feishu list or search endpoint.

        Args:
            method: HTTP method, either ``GET`` or ``POST``.
            path: Feishu OpenAPI path.
            list_key: Key containing page items in the response data.
            body: Optional POST JSON body.

        Returns:
            Concatenated page items.
        """

        items: list[dict[str, Any]] = []
        page_token: str | None = None
        while True:
            params: dict[str, Any] = {"page_size": 500}
            if page_token:
                # Feishu records/search advances pagination through query
                # params. Putting page_token into the POST body repeats page 1.
                params["page_token"] = page_token
            data = self._request(method=method, path=path, params=params, json_body=body)
            items.extend(data.get(list_key, []))
            if not data.get("has_more"):
                return items
            page_token = data.get("page_token")

    def _request(
        self,
        *,
        method: str,
        path: str,
        params: dict[str, Any],
        json_body: dict[str, Any] | None,
    ) -> dict[str, Any]:
        """Send one Feishu OpenAPI request and return the ``data`` payload.

        Args:
            method: HTTP method.
            path: Feishu OpenAPI path.
            params: Query parameters.
            json_body: Optional request JSON body.

        Returns:
            Response ``data`` object.

        Raises:
            RuntimeError: If Feishu returns a non-zero API code.
        """

        response = self._session.request(
            method=method,
            url=f"https://open.feishu.cn{path}",
            headers={
                "Authorization": f"Bearer {self._token_provider.get_token()}",
                "Content-Type": "application/json; charset=utf-8",
            },
            params=params,
            json=json_body,
            timeout=self._timeout_seconds,
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("code") != 0:
            raise RuntimeError(f"Feishu API error code={payload.get('code')}: {payload.get('msg')}")
        data = payload.get("data")
        return data if isinstance(data, dict) else {}


class _TenantTokenProvider:
    """Small tenant-token provider for read-only Feishu grader calls."""

    def __init__(self, *, app_id: str, app_secret: str, timeout_seconds: int, trust_env: bool) -> None:
        """Initialize the token provider.

        Args:
            app_id: Feishu application id.
            app_secret: Feishu application secret.
            timeout_seconds: HTTP timeout in seconds.
            trust_env: Whether requests should inherit host proxy settings.

        Returns:
            None.
        """

        self._app_id = app_id
        self._app_secret = app_secret
        self._timeout_seconds = timeout_seconds
        self._session = requests.Session()
        self._session.trust_env = trust_env
        self._cached_token: str | None = None

    def get_token(self) -> str:
        """Return a tenant access token.

        Args:
            None.

        Returns:
            Feishu tenant access token.
        """

        if self._cached_token:
            return self._cached_token
        response = self._session.post(
            "https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
            json={"app_id": self._app_id, "app_secret": self._app_secret},
            timeout=self._timeout_seconds,
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("code") != 0:
            raise RuntimeError(f"Feishu auth error code={payload.get('code')}: {payload.get('msg')}")
        self._cached_token = str(payload["tenant_access_token"])
        return self._cached_token


def grade_preview_matches_baseline(
    *,
    outputs_dir: Path,
    baseline_source: BaselineSource | None = None,
) -> dict[str, Any]:
    """Return score 1 when preview Excel data matches the Feishu baseline.

    Args:
        outputs_dir: Public case ``outputs/`` directory.
        baseline_source: Optional read-only baseline source. Tests pass a fake
            source; production uses Feishu settings from the environment.

    Returns:
        Grader result for ``result.json.graders[]``.
    """

    preview_file = _find_preview_file(outputs_dir)
    if preview_file is None:
        return _failed_result(
            summary="未在 outputs/ 中找到 preview 产物，无法做 baseline 对账。",
            evidence={"checked_dir": "outputs/"},
        )
    if preview_file.suffix.lower() != ".xlsx":
        return _failed_result(
            summary="baseline 对账 grader 只支持 preview Excel（.xlsx）产物。",
            evidence={"file": _outputs_relative_path(preview_file, outputs_dir), "suffix": preview_file.suffix},
        )

    try:
        preview = read_preview_artifact(preview_file)
        recog_id = _required_metadata_text(preview.metadata, "recog_id")
        period = _required_metadata_text(preview.metadata, "period")
        group_fields = _required_group_fields(preview.metadata)
        key_fields = [PERIOD_FIELD_NAME, *group_fields]
        field_names = _baseline_field_names(preview.headers)
        source = baseline_source or FeishuBaselineSource.from_env()
        table_id = source.table_id_for_recog_id(recog_id)
        if not table_id:
            return _failed_result(
                summary=f"未在项目目录中找到 recog_id={recog_id} 的结果表配置。",
                evidence={
                    "file": _outputs_relative_path(preview_file, outputs_dir),
                    "recog_id": recog_id,
                    "period": period,
                },
            )

        baseline_records = source.fetch_period_records(table_id=table_id, period=period, field_names=field_names)
        report = compare_preview_to_baseline(
            preview_records=preview.records,
            baseline_records=baseline_records,
            key_fields=key_fields,
            compared_fields=field_names,
        )
        evidence = _report_evidence(
            preview_file=preview_file,
            outputs_dir=outputs_dir,
            recog_id=recog_id,
            period=period,
            table_id=table_id,
            report=report,
        )
        if report.matched:
            return {
                "id": "preview_matches_baseline",
                "type": "code",
                "score": 1,
                "summary": (
                    f"preview Excel 与影子飞书 baseline 一致；"
                    f"recog_id={recog_id}，period={period}，row_count={report.preview_count}。"
                ),
                "evidence": evidence,
            }
        return _failed_result(
            summary=(
                f"preview Excel 与影子飞书 baseline 不一致；"
                f"diff_type={_summarize_diff_type(report)}，difference_count={len(report.differences)}。"
            ),
            evidence=evidence,
        )
    except GraderConfigurationError as exc:
        return _failed_result(
            summary="baseline 对账所需环境变量缺失。",
            evidence={"file": _outputs_relative_path(preview_file, outputs_dir), "error": str(exc)},
        )
    except Exception as exc:  # noqa: BLE001
        return _failed_result(
            summary="baseline 对账 grader 执行失败。",
            evidence={
                "file": _outputs_relative_path(preview_file, outputs_dir),
                "error_type": type(exc).__name__,
                "error": str(exc),
            },
        )


def read_preview_artifact(result_file: Path) -> PreviewArtifactData:
    """Read records and metadata from a preview Excel artifact.

    Args:
        result_file: Preview Excel path.

    Returns:
        PreviewArtifactData containing result rows and hidden metadata.
    """

    workbook = load_workbook(result_file, data_only=True, read_only=True)
    try:
        if PREVIEW_RESULTS_SHEET not in workbook.sheetnames:
            raise ValueError(f"preview Excel missing sheet: {PREVIEW_RESULTS_SHEET}")
        if PREVIEW_METADATA_SHEET not in workbook.sheetnames:
            raise ValueError(f"preview Excel missing sheet: {PREVIEW_METADATA_SHEET}")

        results_sheet = workbook[PREVIEW_RESULTS_SHEET]
        rows = list(results_sheet.iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in rows[0]] if rows else []
        records = [
            {headers[index]: value for index, value in enumerate(row[: len(headers)]) if headers[index]}
            for row in rows[1:]
            if any(value is not None for value in row)
        ]

        metadata_sheet = workbook[PREVIEW_METADATA_SHEET]
        metadata: dict[str, Any] = {}
        for key_cell, value_cell in metadata_sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            if key_cell is None:
                continue
            key = str(key_cell)
            if key == "group_fields" and isinstance(value_cell, str):
                metadata[key] = json.loads(value_cell)
            else:
                metadata[key] = value_cell
        return PreviewArtifactData(records=records, metadata=metadata, headers=[header for header in headers if header])
    finally:
        workbook.close()


def compare_preview_to_baseline(
    *,
    preview_records: list[dict[str, Any]],
    baseline_records: list[dict[str, Any]],
    key_fields: list[str],
    compared_fields: list[str],
    numeric_digits: int = 2,
) -> BaselineDiffReport:
    """Compare preview rows with baseline rows by business key.

    Args:
        preview_records: Preview rows from Excel.
        baseline_records: Feishu baseline rows.
        key_fields: Row alignment fields, usually ``期间`` plus group fields.
        compared_fields: Field names to compare.
        numeric_digits: Decimal places used for numeric comparison.

    Returns:
        BaselineDiffReport with all detected differences.
    """

    preview_map = _index_records(records=preview_records, key_fields=key_fields, numeric_digits=numeric_digits)
    baseline_map = _index_records(records=baseline_records, key_fields=key_fields, numeric_digits=numeric_digits)
    differences: list[DifferenceEntry] = []

    for key in sorted(set(preview_map) | set(baseline_map), key=str):
        preview_record = preview_map.get(key)
        baseline_record = baseline_map.get(key)
        if preview_record is None:
            differences.append(
                DifferenceEntry(
                    key=key,
                    issue_type="missing_in_preview",
                    differences=[DifferenceDetail(field="__record__", preview_value=None, baseline_value=baseline_record)],
                )
            )
            continue
        if baseline_record is None:
            differences.append(
                DifferenceEntry(
                    key=key,
                    issue_type="missing_in_baseline",
                    differences=[DifferenceDetail(field="__record__", preview_value=preview_record, baseline_value=None)],
                )
            )
            continue

        field_differences = [
            DifferenceDetail(
                field=field_name,
                preview_value=preview_record.get(field_name),
                baseline_value=baseline_record.get(field_name),
            )
            for field_name in compared_fields
            if field_name not in key_fields and preview_record.get(field_name) != baseline_record.get(field_name)
        ]
        if field_differences:
            differences.append(DifferenceEntry(key=key, issue_type="value_mismatch", differences=field_differences))

    return BaselineDiffReport(
        matched=not differences,
        preview_count=len(preview_records),
        baseline_count=len(baseline_records),
        key_fields=key_fields,
        compared_fields=compared_fields,
        differences=differences,
    )


def flatten_feishu_fields(fields: dict[str, Any]) -> dict[str, Any]:
    """Flatten Feishu field payloads to plain comparable values.

    Args:
        fields: Feishu ``record.fields`` payload.

    Returns:
        Plain dictionary with scalar values where possible.
    """

    return {field_name: extract_feishu_scalar(field_value) for field_name, field_value in fields.items()}


def extract_feishu_scalar(value: Any) -> Any:
    """Extract a comparable scalar from common Feishu field value shapes.

    Args:
        value: Raw Feishu field value.

    Returns:
        Scalar value, list of scalar values, or nested dictionary.
    """

    if isinstance(value, list):
        if not value:
            return None
        if len(value) == 1:
            return extract_feishu_scalar(value[0])
        return [extract_feishu_scalar(item) for item in value]
    if isinstance(value, dict):
        if "text" in value:
            return value["text"]
        if "type" in value and "value" in value:
            return extract_feishu_scalar(value["value"])
        if "value" in value:
            return extract_feishu_scalar(value["value"])
        return {key: extract_feishu_scalar(inner_value) for key, inner_value in value.items()}
    return value


def normalize_scalar(value: Any, *, digits: int = 2) -> Any:
    """Normalize one non-key field value for comparison.

    Args:
        value: Raw field value.
        digits: Decimal places used for numeric comparison.

    Returns:
        Comparable normalized value.
    """

    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        if _is_dash_placeholder(stripped):
            return 0.0
        parsed = _try_parse_decimal(stripped)
        if parsed is not None:
            return _quantize_decimal(parsed, digits=digits)
        return normalize_display_text(stripped)
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)):
        return _quantize_decimal(Decimal(str(value)), digits=digits)
    return value


def normalize_display_text(value: str) -> str:
    """Normalize low-risk display text differences.

    Args:
        value: Raw display text.

    Returns:
        Text with full-width parentheses unified to half-width parentheses.
    """

    return value.translate(str.maketrans({"（": "(", "）": ")"}))


def _failed_result(*, summary: str, evidence: dict[str, Any]) -> dict[str, Any]:
    """Build a failed grader result.

    Args:
        summary: Human-readable failure summary.
        evidence: Structured evidence safe to write into ``result.json``.

    Returns:
        Grader result dictionary.
    """

    return {
        "id": "preview_matches_baseline",
        "type": "code",
        "score": 0,
        "summary": summary,
        "evidence": evidence,
    }


def _report_evidence(
    *,
    preview_file: Path,
    outputs_dir: Path,
    recog_id: str,
    period: str,
    table_id: str,
    report: BaselineDiffReport,
) -> dict[str, Any]:
    """Convert a diff report to compact grader evidence.

    Args:
        preview_file: Preview artifact path.
        outputs_dir: Public case outputs directory.
        recog_id: Recognition project id.
        period: Accounting period.
        table_id: Feishu result table id.
        report: Full comparison report.

    Returns:
        JSON-serializable evidence dictionary.
    """

    return {
        "file": _outputs_relative_path(preview_file, outputs_dir),
        "recog_id": recog_id,
        "period": period,
        "target_table_id": table_id,
        "preview_count": report.preview_count,
        "baseline_count": report.baseline_count,
        "key_fields": report.key_fields,
        "compared_fields": report.compared_fields,
        "difference_count": len(report.differences),
        "diff_type": _summarize_diff_type(report),
        "sample_differences": [_difference_to_evidence(item) for item in report.differences[:MAX_SAMPLE_DIFFERENCES]],
    }


def _difference_to_evidence(difference: DifferenceEntry) -> dict[str, Any]:
    """Convert one difference entry to compact JSON evidence.

    Args:
        difference: Difference entry.

    Returns:
        JSON-serializable difference summary.
    """

    return {
        "key": list(difference.key),
        "issue_type": difference.issue_type,
        "differences": [
            {
                "field": detail.field,
                "preview_value": detail.preview_value,
                "baseline_value": detail.baseline_value,
            }
            for detail in difference.differences[:MAX_SAMPLE_FIELD_DIFFERENCES]
        ],
    }


def _summarize_diff_type(report: BaselineDiffReport) -> str:
    """Summarize a diff report into one stable category.

    Args:
        report: Comparison report.

    Returns:
        ``matched`` or a compact issue category.
    """

    if report.matched:
        return "matched"
    issue_types = {difference.issue_type for difference in report.differences}
    if len(issue_types) == 1:
        return next(iter(issue_types))
    return "mixed"


def _required_metadata_text(metadata: dict[str, Any], key: str) -> str:
    """Read one required metadata value as stripped text.

    Args:
        metadata: Preview metadata dictionary.
        key: Required metadata key.

    Returns:
        Non-empty string value.

    Raises:
        ValueError: If the key is missing or empty.
    """

    value = metadata.get(key)
    normalized = str(value).strip() if value is not None else ""
    if not normalized:
        raise ValueError(f"preview metadata missing required key: {key}")
    return normalized


def _required_group_fields(metadata: dict[str, Any]) -> list[str]:
    """Read required preview group fields from metadata.

    Args:
        metadata: Preview metadata dictionary.

    Returns:
        Non-empty group field list.

    Raises:
        ValueError: If group fields are missing or invalid.
    """

    value = metadata.get("group_fields")
    if not isinstance(value, list) or not all(isinstance(item, str) and item.strip() for item in value):
        raise ValueError("preview metadata missing valid group_fields")
    return [item.strip() for item in value]


def _baseline_field_names(headers: list[str]) -> list[str]:
    """Build the field list requested from Feishu baseline.

    Args:
        headers: Preview result-sheet headers.

    Returns:
        Ordered unique field names, always including ``期间``.
    """

    ordered: list[str] = []
    for field_name in [PERIOD_FIELD_NAME, *headers]:
        if field_name and field_name not in ordered:
            ordered.append(field_name)
    return ordered


def _index_records(
    *,
    records: list[dict[str, Any]],
    key_fields: list[str],
    numeric_digits: int,
) -> dict[tuple[Any, ...], dict[str, Any]]:
    """Index records by normalized business key.

    Args:
        records: Raw record dictionaries.
        key_fields: Field names forming the business key.
        numeric_digits: Decimal places used for non-key numeric values.

    Returns:
        Mapping from business key to normalized record.
    """

    indexed: dict[tuple[Any, ...], dict[str, Any]] = {}
    for record in records:
        key = tuple(_normalize_key_value(record.get(field_name)) for field_name in key_fields)
        indexed[key] = {field_name: normalize_scalar(value, digits=numeric_digits) for field_name, value in record.items()}
    return indexed


def _normalize_key_value(value: Any) -> Any:
    """Normalize business-key values without losing text id precision.

    Args:
        value: Raw key value.

    Returns:
        Comparable key value.
    """

    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return normalize_display_text(stripped) if stripped else None
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        return str(int(value)) if value == value.to_integral_value() else format(value, "f")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return normalize_display_text(str(value).strip())


def _extract_feishu_text(value: Any) -> str | None:
    """Extract a text value from common Feishu catalog field shapes.

    Args:
        value: Raw Feishu field value.

    Returns:
        Text value, or ``None`` when empty.
    """

    scalar = extract_feishu_scalar(value)
    if scalar is None:
        return None
    if isinstance(scalar, list):
        for item in scalar:
            extracted = _extract_feishu_text(item)
            if extracted:
                return extracted
        return None
    text = str(scalar).strip()
    return text or None


def _ensure_period_field(field_names: list[str]) -> list[str]:
    """Ensure the Feishu field list contains ``期间``.

    Args:
        field_names: Requested field names.

    Returns:
        Ordered unique field names including ``期间``.
    """

    return _baseline_field_names(field_names)


def _build_period_filter_value(period: str, field_type: int | None) -> int | str:
    """Build a Feishu filter value that matches the remote period field type.

    Args:
        period: Accounting period in ``YYYYMM`` format.
        field_type: Feishu field type id.

    Returns:
        Integer period for numeric fields; string period otherwise.
    """

    normalized = str(period).strip()
    if len(normalized) != 6 or not normalized.isdigit():
        raise ValueError("period must be a 6-digit YYYYMM token")
    if field_type in NUMERIC_FIELD_TYPES:
        return int(normalized)
    if field_type in TEXTUAL_FIELD_TYPES:
        return normalized
    return normalized


def _try_parse_decimal(value: str) -> Decimal | None:
    """Parse a human-formatted number into Decimal when possible.

    Args:
        value: Candidate string.

    Returns:
        Parsed Decimal or ``None``.
    """

    candidate = value.replace(",", "")
    try:
        return Decimal(candidate)
    except (InvalidOperation, ValueError):
        return None


def _quantize_decimal(value: Decimal, *, digits: int) -> float:
    """Round a Decimal value using the baseline comparison precision.

    Args:
        value: Decimal value.
        digits: Decimal places.

    Returns:
        Float value rounded with ``ROUND_HALF_UP``.
    """

    quantizer = Decimal("1." + "0" * digits)
    return float(value.quantize(quantizer, rounding=ROUND_HALF_UP))


def _is_dash_placeholder(value: str) -> bool:
    """Return whether a string is only dash-like placeholder characters.

    Args:
        value: Candidate text.

    Returns:
        True when the text is made only of dash placeholders.
    """

    return bool(re.fullmatch(r"[-—–－]+", value))
