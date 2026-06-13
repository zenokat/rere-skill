"""preview 结果文件写入器。"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from models.domain import PREVIEW_METADATA_SHEET, PREVIEW_RESULTS_SHEET, PreviewArtifact
from utils.output_paths import build_preview_output_path, ensure_output_root


class PreviewArtifactWriter:
    """负责写入和读取 preview 结果文件。"""

    def __init__(self, output_root: Path) -> None:
        """初始化输出目录。"""

        self._output_root = output_root

    def write(
        self,
        recog_id: str,
        period: str,
        records: list[dict[str, Any]],
        group_fields: list[str],
    ) -> PreviewArtifact:
        """把试算结果写入 Excel 文件。"""

        ensure_output_root(self._output_root)
        generated_at = datetime.now()
        result_file = build_preview_output_path(period=period, recog_id=recog_id, output_root=self._output_root, generated_at=generated_at)

        workbook = Workbook()
        results_sheet = workbook.active
        results_sheet.title = PREVIEW_RESULTS_SHEET

        headers = list(records[0].keys()) if records else self._build_default_headers(group_fields=group_fields)
        results_sheet.append(headers)
        for record in records:
            results_sheet.append([record.get(header) for header in headers])

        metadata_sheet = workbook.create_sheet(PREVIEW_METADATA_SHEET)
        metadata_sheet.sheet_state = "hidden"
        metadata_sheet.append(["key", "value"])
        metadata_sheet.append(["recog_id", recog_id])
        metadata_sheet.append(["period", period])
        metadata_sheet.append(["generated_at", generated_at.isoformat()])
        metadata_sheet.append(["group_fields", json.dumps(group_fields, ensure_ascii=False)])
        workbook.save(result_file)
        workbook.close()

        field_count = len(headers)
        row_count = len(records)
        return PreviewArtifact(
            recog_id=recog_id,
            period=period,
            result_file=result_file,
            row_count=row_count,
            field_count=field_count,
            generated_at=generated_at.isoformat(),
            group_fields=group_fields,
        )

    def read(self, result_file: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        """读取 preview 结果文件及元数据。"""

        workbook = load_workbook(result_file, data_only=True)
        try:
            results_sheet = workbook[PREVIEW_RESULTS_SHEET]
            metadata_sheet = workbook[PREVIEW_METADATA_SHEET]

            rows = list(results_sheet.iter_rows(values_only=True))
            headers = [str(header) for header in rows[0]] if rows else []
            records = [
                {headers[index]: value for index, value in enumerate(row)}
                for row in rows[1:]
                if headers
            ]

            metadata: dict[str, Any] = {}
            for key_cell, value_cell in metadata_sheet.iter_rows(min_row=2, max_col=2, values_only=True):
                if key_cell == "group_fields" and isinstance(value_cell, str):
                    metadata[str(key_cell)] = json.loads(value_cell)
                else:
                    metadata[str(key_cell)] = value_cell
            return records, metadata
        finally:
            workbook.close()

    @staticmethod
    def _build_default_headers(group_fields: list[str]) -> list[str]:
        """在无结果记录时保底生成表头。"""

        return ["期间", *group_fields]
