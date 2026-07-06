"""源文件解析缓存的单元测试。"""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook

from core import source_resolver
from models.domain import SourceSheetSpec

TMP_ROOT = Path.cwd() / "tests" / "_tmp_source_resolver"


def _make_tmp_dir() -> Path:
    """在仓库内创建可写临时目录，避免系统临时目录权限问题。"""

    tmp_dir = TMP_ROOT / uuid4().hex
    tmp_dir.mkdir(parents=True, exist_ok=True)
    return tmp_dir


def test_build_sheet_dataset_reuses_workbook_snapshot_for_same_file(monkeypatch) -> None:
    """同一工作簿解析多个 sheet 时应复用一次加载结果。"""

    target_file = Path("demo.xlsx")
    load_cache = source_resolver.SourceFileLoadCache()
    snapshot_calls: list[Path] = []

    def fake_build_workbook_snapshot(source_file: Path) -> source_resolver.WorkbookSnapshot:
        """返回固定快照，并记录被调用次数。"""

        snapshot_calls.append(source_file)
        return source_resolver.WorkbookSnapshot(
            sheet_names=("SheetA", "SheetB"),
            raw_rows_by_sheet={
                "SheetA": [["门店", "金额"], ["A店", "10"]],
                "SheetB": [["门店", "金额"], ["B店", "20"]],
            },
        )

    monkeypatch.setattr(source_resolver, "_build_workbook_snapshot", fake_build_workbook_snapshot)

    source_resolver.build_sheet_dataset(
        target_file,
        SourceSheetSpec(recog_id="demo", sheet="SheetA", field_row=1, last_row=-1),
        load_cache=load_cache,
    )
    source_resolver.build_sheet_dataset(
        target_file,
        SourceSheetSpec(recog_id="demo", sheet="SheetB", field_row=1, last_row=-1),
        load_cache=load_cache,
    )

    assert snapshot_calls == [target_file]


def test_build_sheet_dataset_injects_file_name_virtual_column() -> None:
    """CSV 数据集应自动提供 `FILE_NAME` 虚拟列。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "wechat-demo.csv"
    source_file.write_text(
        "门店,金额\n"
        "A店,10\n",
        encoding="utf-8",
    )

    dataset = source_resolver.build_sheet_dataset(
        source_file,
        SourceSheetSpec(recog_id="demo", sheet="DEFAULT", field_row=1, last_row=-1),
    )

    assert any(column.field == "FILE_NAME" for column in dataset.columns)
    assert dataset.data_rows[0][("", "FILE_NAME")] == "wechat-demo"


def test_build_sheet_dataset_strips_backtick_prefix_from_csv_values() -> None:
    """CSV 单元格中的前缀反引号应被清理。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "wechat-demo.csv"
    source_file.write_text(
        "平台ID,平台门店名称,订单金额\n"
        "`1001,`蔡林记中南店,`12.50\n",
        encoding="utf-8",
    )

    dataset = source_resolver.build_sheet_dataset(
        source_file,
        SourceSheetSpec(recog_id="demo", sheet="DEFAULT", field_row=1, last_row=-1),
    )

    row = dataset.data_rows[0]
    assert row[("", "平台ID")] == "1001"
    assert row[("", "平台门店名称")] == "蔡林记中南店"
    assert row[("", "订单金额")] == "12.50"


def test_build_sheet_dataset_skips_blank_csv_rows() -> None:
    """CSV 中的全空白行不应占用 `field_row` 计数。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "alipay-demo.csv"
    source_file.write_text(
        "#标题\n"
        "\n"
        "#账号\n"
        "\n"
        "#日期\n"
        "\n"
        "#列表\n"
        "\n"
        "门店编号,商家实收（元）,服务费（元）\n"
        "1001,12.50,-0.03\n",
        encoding="utf-8",
    )

    dataset = source_resolver.build_sheet_dataset(
        source_file,
        SourceSheetSpec(recog_id="demo", sheet="DEFAULT", field_row=5, last_row=-1),
    )

    fields = [column.field for column in dataset.columns]
    assert "门店编号" in fields
    assert "商家实收（元）" in fields
    assert dataset.data_rows[0][("", "门店编号")] == "1001"


def test_build_sheet_dataset_reads_gb18030_csv() -> None:
    """常见中文账单 CSV 即使不是 UTF-8，也应能被稳定读取。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "meituan-demo.csv"
    source_file.write_text(
        "日期,门店名称,商品原价\n"
        "2026-05-01,蔡林记,12.50\n",
        encoding="gb18030",
    )

    dataset = source_resolver.build_sheet_dataset(
        source_file,
        SourceSheetSpec(recog_id="demo", sheet="DEFAULT", field_row=1, last_row=-1),
    )

    row = dataset.data_rows[0]
    assert row[("", "日期")] == "2026-05-01"
    assert row[("", "门店名称")] == "蔡林记"
    assert row[("", "商品原价")] == "12.50"


def test_streaming_sheet_dataset_skips_blank_excel_rows() -> None:
    """大 Excel 的流式读取应跳过尾部全空白数据行。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "meituan-summary.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "按门店"
    worksheet.append(["门店名称", "服务费"])
    worksheet.append(["蔡林记", "12.50"])
    worksheet.append([None, None])
    worksheet.append([None, None])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="按门店", field_row=1, last_row=-1),
            stream_to_end=True,
        )
        rows = list(dataset.iter_data_rows())
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert len(rows) == 1
    assert rows[0][0] == 2
    assert rows[0][1][("", "门店名称")] == "蔡林记"
    assert rows[0][1][("", "服务费")] == "12.50"


def test_non_streaming_sheet_dataset_skips_blank_excel_rows() -> None:
    """普通 xlsx 解析路径也应跳过尾部全空白数据行。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "meituan-summary-small.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "按门店"
    worksheet.append(["门店名称", "服务费"])
    worksheet.append(["蔡林记", "12.50"])
    worksheet.append([None, None])
    worksheet.append([None, None])
    workbook.save(source_file)
    workbook.close()

    dataset = source_resolver.build_sheet_dataset(
        source_file,
        SourceSheetSpec(recog_id="demo", sheet="按门店", field_row=1, last_row=-1),
    )

    assert dataset.data_row_numbers == [2]
    assert len(dataset.data_rows) == 1
    assert dataset.data_rows[0][("", "门店名称")] == "蔡林记"
    assert dataset.data_rows[0][("", "服务费")] == "12.50"


def test_streaming_sheet_dataset_allows_header_only_excel_when_last_row_is_to_end() -> None:
    """只有表头、没有数据体的大 Excel 分片应被视为零行数据集，而不是报空区间错误。"""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "jingdong-header-only.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "com.jd.o2o.settlement.domain.dt"
    worksheet.append(["商家基础信息", None, None])
    worksheet.append(["商家编号", "商家名称", "门店编号"])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="com.jd.o2o.settlement.domain.dt", field_row=2, last_row=-1),
            stream_to_end=True,
        )
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert dataset.first_data_row is None
    assert dataset.first_data_row_number is None
    assert list(dataset.iter_data_rows()) == []


def test_streaming_sheet_dataset_supports_category_header_rows() -> None:
    """Large Excel streaming should support two-row category and field headers."""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "category-header.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet.append(["store", "settlement", "settlement"])
    worksheet.append(["name", "amount", "fee"])
    worksheet.append(["Shop A", "12.50", "-0.30"])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="Summary", category_row=1, field_row=2, last_row=-1),
            stream_to_end=True,
        )
        rows = list(dataset.iter_data_rows())
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert rows[0][0] == 3
    assert rows[0][1][("store", "name")] == "Shop A"
    assert rows[0][1][("settlement", "amount")] == "12.50"
    assert rows[0][1][("settlement", "fee")] == "-0.30"


def test_streaming_sheet_dataset_falls_back_for_merged_category_headers() -> None:
    """Streaming should not lose categories that depend on merged Excel cells."""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "merged-category-header.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet["A1"] = "settlement"
    worksheet.merge_cells("A1:B1")
    worksheet.append(["amount", "fee"])
    worksheet.append(["12.50", "-0.30"])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="Summary", category_row=1, field_row=2, last_row=-1),
            stream_to_end=True,
        )
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert dataset.row_iterator_factory is None
    assert dataset.data_rows[0][("settlement", "amount")] == "12.50"
    assert dataset.data_rows[0][("settlement", "fee")] == "-0.30"


def test_streaming_sheet_dataset_reports_progress_for_rows() -> None:
    """Streaming row iteration should report physical row numbers for progress logs."""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "progress.xlsx"
    reported_rows: list[int] = []

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rows"
    worksheet.append(["name", "amount"])
    worksheet.append(["Shop A", "12.50"])
    worksheet.append(["Shop B", "18.00"])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="Rows", field_row=1, last_row=-1),
            stream_to_end=True,
            progress_callback=reported_rows.append,
        )
        rows = list(dataset.iter_data_rows())
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert [row_number for row_number, _ in rows] == [2, 3]
    assert reported_rows == [2, 3]


def test_streaming_sheet_dataset_reports_progress_for_fixed_last_row() -> None:
    """Streaming with an explicit last row should use the same progress callback."""

    tmp_dir = _make_tmp_dir()
    source_file = tmp_dir / "fixed-range-progress.xlsx"
    reported_rows: list[int] = []

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rows"
    worksheet.append(["name", "amount"])
    worksheet.append(["Shop A", "12.50"])
    worksheet.append(["Shop B", "18.00"])
    worksheet.append(["Ignored", "99.00"])
    workbook.save(source_file)
    workbook.close()

    original_threshold = source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD
    source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = 0
    try:
        dataset = source_resolver.build_sheet_dataset(
            source_file,
            SourceSheetSpec(recog_id="demo", sheet="Rows", field_row=1, last_row=3),
            progress_callback=reported_rows.append,
        )
        rows = list(dataset.iter_data_rows())
    finally:
        source_resolver.STREAMING_WORKBOOK_SIZE_THRESHOLD = original_threshold

    assert [row_number for row_number, _ in rows] == [2, 3]
    assert reported_rows == [2, 3]


def test_prepare_read_only_worksheet_resets_stale_dimensions_even_when_max_row_is_not_one() -> None:
    """只读 worksheet 即使暴露出的 max_row 不是 1，也应主动重算失真的 dimension。"""

    class _FakeWorksheet:
        def __init__(self) -> None:
            self.max_row = 8
            self.reset_calls = 0
            self.calculate_calls: list[bool] = []

        def reset_dimensions(self) -> None:
            self.reset_calls += 1

        def calculate_dimension(self, force: bool = False) -> str:
            self.calculate_calls.append(force)
            return "A1:BH162747"

    worksheet = _FakeWorksheet()

    source_resolver._prepare_read_only_worksheet(worksheet, force_calculate=True)

    assert worksheet.reset_calls == 1
    assert worksheet.calculate_calls == [True]
