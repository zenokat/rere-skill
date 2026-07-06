"""源文件发现与逻辑 sheet 解析工具。

本模块负责解决两类问题：
1. `source_file` 既可能是单文件，也可能是目录
2. `csv` 或单 sheet `xlsx` 需要支持 `DEFAULT` 逻辑 sheet
"""

from __future__ import annotations

import csv
import random
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from models.cli_results import CliExecutionError, ErrorResponse, ExitCode
from models.domain import DEFAULT_LOGICAL_SHEET, SourceSheetSpec

SUPPORTED_SOURCE_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
STREAMING_WORKBOOK_SIZE_THRESHOLD = 10 * 1024 * 1024
DataRow = dict[tuple[str, str], Any]
DataRowIteratorFactory = Callable[[], Iterator[tuple[int, DataRow]]]
ProgressCallback = Callable[[int], None]
TEMPORARY_SOURCE_PREFIXES = ("~$",)
FILE_NAME_VIRTUAL_FIELD = "FILE_NAME"
SUPPORTED_CSV_ENCODINGS = ("utf-8-sig", "gb18030")


@dataclass(frozen=True)
class ColumnDescriptor:
    """描述一个可被规则引用的源列。"""

    column_index: int
    category: str | None
    field: str


@dataclass
class SheetDataset:
    """一个物理文件中某个逻辑 sheet 的解析结果。"""

    source_file: Path
    logical_sheet: str
    physical_sheet: str
    total_rows: int
    resolved_last_row: int
    columns: list[ColumnDescriptor]
    duplicate_columns: dict[tuple[str, str], list[int]]
    data_rows: list[DataRow]
    data_row_numbers: list[int]
    first_data_row: DataRow | None = None
    first_data_row_number: int | None = None
    row_iterator_factory: DataRowIteratorFactory | None = None

    def iter_data_rows(self) -> Iterator[tuple[int, DataRow]]:
        """按统一接口遍历数据行。"""

        if self.row_iterator_factory is not None:
            yield from self.row_iterator_factory()
            return

        for row_number, row_values in zip(self.data_row_numbers, self.data_rows):
            yield row_number, row_values

    def sample_row_values(self) -> DataRow:
        """返回一行代表性数据，供 condition 预编译使用。"""

        if self.first_data_row is not None:
            return self.first_data_row
        if self.data_rows:
            return self.data_rows[0]
        return {}


@dataclass
class WorkbookSnapshot:
    """缓存单个工作簿全部 sheet 的原始行数据。"""

    sheet_names: tuple[str, ...]
    raw_rows_by_sheet: dict[str, list[list[Any]]]

    def resolve_rows(self, logical_sheet: str) -> tuple[list[list[Any]], str]:
        """按逻辑 sheet 读取缓存的原始行数据。"""

        if logical_sheet == DEFAULT_LOGICAL_SHEET:
            if len(self.sheet_names) != 1:
                raise ValueError("DEFAULT logical sheet can only be used when the workbook has exactly one sheet.")
            physical_sheet = self.sheet_names[0]
        else:
            if logical_sheet not in self.raw_rows_by_sheet:
                raise KeyError(logical_sheet)
            physical_sheet = logical_sheet
        return self.raw_rows_by_sheet[physical_sheet], physical_sheet


@dataclass
class SourceFileLoadCache:
    """在一次运行周期内复用源文件解析结果。

    对于多 sheet 的大型工作簿，preview/validate 会在同一文件上反复读取多个 sheet。
    这里把工作簿内容按“单文件一次加载”的方式缓存下来，避免重复打开同一个 Excel 文件。
    """

    csv_rows_by_file: dict[Path, list[list[Any]]] = field(default_factory=dict)
    workbook_snapshots_by_file: dict[Path, WorkbookSnapshot] = field(default_factory=dict)

    def get_raw_rows(self, source_file: Path, logical_sheet: str) -> tuple[list[list[Any]], str]:
        """返回某个逻辑 sheet 对应的原始二维表格。"""

        suffix = source_file.suffix.lower()
        if suffix == ".csv":
            return self._get_csv_rows(source_file=source_file, logical_sheet=logical_sheet)
        return self._get_workbook_rows(source_file=source_file, logical_sheet=logical_sheet)

    def _get_csv_rows(self, source_file: Path, logical_sheet: str) -> tuple[list[list[Any]], str]:
        """读取并缓存 CSV 原始行数据。"""

        if logical_sheet != DEFAULT_LOGICAL_SHEET:
            raise ValueError("CSV files must use the DEFAULT logical sheet.")
        raw_rows = self.csv_rows_by_file.get(source_file)
        if raw_rows is None:
            raw_rows = _read_csv_rows(source_file)
            self.csv_rows_by_file[source_file] = raw_rows
        return raw_rows, DEFAULT_LOGICAL_SHEET

    def _get_workbook_rows(self, source_file: Path, logical_sheet: str) -> tuple[list[list[Any]], str]:
        """读取并缓存工作簿快照。"""

        snapshot = self.workbook_snapshots_by_file.get(source_file)
        if snapshot is None:
            snapshot = _build_workbook_snapshot(source_file)
            self.workbook_snapshots_by_file[source_file] = snapshot
        return snapshot.resolve_rows(logical_sheet)


def discover_source_files(source_path: Path) -> list[Path]:
    """根据单文件或目录路径找出全部待处理源文件。"""

    if not source_path.exists():
        raise CliExecutionError(
            exit_code=ExitCode.LOCAL_FILE_FAILED,
            response=ErrorResponse(
                stage="source_file",
                message="The source_file path does not exist.",
                retryable=False,
                details=[{"source_file": str(source_path)}],
            ),
        )

    if source_path.is_file():
        if _should_skip_source_file(source_path):
            raise CliExecutionError(
                exit_code=ExitCode.LOCAL_FILE_FAILED,
                response=ErrorResponse(
                    stage="source_file",
                    message="The source_file path points to a temporary Office file that should not be processed.",
                    retryable=False,
                    details=[{"source_file": str(source_path)}],
                ),
            )
        if source_path.suffix.lower() not in SUPPORTED_SOURCE_SUFFIXES:
            raise CliExecutionError(
                exit_code=ExitCode.LOCAL_FILE_FAILED,
                response=ErrorResponse(
                    stage="source_file",
                    message="The source file type is not supported.",
                    retryable=False,
                    details=[{"source_file": str(source_path), "suffix": source_path.suffix.lower()}],
                ),
            )
        return [source_path]

    files = sorted(
        file_path
        for file_path in source_path.rglob("*")
        if file_path.is_file()
        and file_path.suffix.lower() in SUPPORTED_SOURCE_SUFFIXES
        and not _should_skip_source_file(file_path)
    )
    if not files:
        raise CliExecutionError(
            exit_code=ExitCode.LOCAL_FILE_FAILED,
            response=ErrorResponse(
                stage="source_file",
                message="No supported source files were found under the source directory.",
                retryable=False,
                details=[{"source_file": str(source_path)}],
            ),
        )

    suffixes = {file_path.suffix.lower() for file_path in files}
    if len(suffixes) > 1:
        raise CliExecutionError(
            exit_code=ExitCode.LOCAL_FILE_FAILED,
            response=ErrorResponse(
                stage="source_file",
                message="The source directory contains mixed file types. Please let the Agent normalize the directory first.",
                retryable=False,
                details=[{"source_file": str(source_path), "suffixes": sorted(suffixes)}],
            ),
        )
    return files


def _should_skip_source_file(source_file: Path) -> bool:
    """判断某个文件是否应从源文件发现中排除。"""

    return source_file.name.startswith(TEMPORARY_SOURCE_PREFIXES)


def choose_validation_sample(source_files: list[Path], seed_key: str) -> Path:
    """从全部源文件中挑选一个代表文件用于 validate。"""

    chooser = random.Random(seed_key)
    return chooser.choice(source_files)


def build_sheet_dataset(
    source_file: Path,
    spec: SourceSheetSpec,
    *,
    load_cache: SourceFileLoadCache | None = None,
    stream_to_end: bool = False,
    progress_callback: ProgressCallback | None = None,
) -> SheetDataset:
    """按源结构定义读取并解析单个逻辑 sheet。"""

    if _should_stream_sheet_dataset(source_file=source_file, spec=spec):
        if stream_to_end and spec.last_row == -1:
            return _build_streaming_sheet_dataset_to_end(
                source_file=source_file,
                spec=spec,
                load_cache=load_cache,
                progress_callback=progress_callback,
            )
        return _build_streaming_sheet_dataset(
            source_file=source_file,
            spec=spec,
            load_cache=load_cache,
            progress_callback=progress_callback,
        )

    return _build_in_memory_sheet_dataset(
        source_file=source_file,
        spec=spec,
        load_cache=load_cache,
    )


def _build_in_memory_sheet_dataset(
    source_file: Path,
    spec: SourceSheetSpec,
    *,
    load_cache: SourceFileLoadCache | None,
) -> SheetDataset:
    """Build a sheet dataset by loading the needed workbook rows into memory.

    Args:
        source_file: CSV or workbook file to parse.
        spec: Source sheet structure from the rule bundle.
        load_cache: Optional per-run cache used to avoid reopening the same
            workbook repeatedly.

    Returns:
        Parsed sheet dataset with data rows materialized in memory.
    """

    raw_rows, physical_sheet = _load_raw_rows(
        source_file=source_file,
        logical_sheet=spec.sheet,
        load_cache=load_cache,
    )
    total_rows = len(raw_rows)
    if total_rows == 0:
        raise ValueError(f"No rows were found in {source_file}.")

    resolved_last_row = _resolve_last_row(total_rows=total_rows, configured_last_row=spec.last_row)
    header_anchor_row = max(spec.category_row or 0, spec.field_row)
    data_start_row = header_anchor_row + 1

    _validate_sheet_bounds(
        spec=spec,
        total_rows=total_rows,
        resolved_last_row=resolved_last_row,
        data_start_row=data_start_row,
        allow_header_only_empty=spec.last_row == -1,
    )

    columns, duplicate_columns = _build_column_descriptors(raw_rows=raw_rows, spec=spec)

    data_rows: list[DataRow] = []
    data_row_numbers: list[int] = []
    for row_number in range(data_start_row, resolved_last_row + 1):
        row_values = _build_row_value_map(
            row_values=raw_rows[row_number - 1],
            columns=columns,
            source_file=source_file,
        )
        if _is_effectively_blank_data_row(row_values):
            continue
        data_rows.append(row_values)
        data_row_numbers.append(row_number)

    return SheetDataset(
        source_file=source_file,
        logical_sheet=spec.sheet,
        physical_sheet=physical_sheet,
        total_rows=total_rows,
        resolved_last_row=resolved_last_row,
        columns=columns,
        duplicate_columns=duplicate_columns,
        data_rows=data_rows,
        data_row_numbers=data_row_numbers,
        first_data_row=data_rows[0] if data_rows else None,
        first_data_row_number=data_row_numbers[0] if data_row_numbers else None,
    )


def _should_stream_sheet_dataset(source_file: Path, spec: SourceSheetSpec) -> bool:
    """判断当前 sheet 是否适合走流式读取路径。"""

    if source_file.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    try:
        return source_file.stat().st_size >= STREAMING_WORKBOOK_SIZE_THRESHOLD
    except OSError:
        return False


def _build_streaming_sheet_dataset(
    source_file: Path,
    spec: SourceSheetSpec,
    *,
    load_cache: SourceFileLoadCache | None = None,
    progress_callback: ProgressCallback | None = None,
) -> SheetDataset:
    """为超大工作簿构建轻量级数据集描述。"""

    from openpyxl import load_workbook

    should_fallback_to_memory = False
    workbook = load_workbook(source_file, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet, physical_sheet = _resolve_workbook_worksheet(workbook=workbook, logical_sheet=spec.sheet)
        _prepare_read_only_worksheet(worksheet, force_calculate=True)
        total_rows = int(worksheet.max_row or _count_worksheet_rows(worksheet))
        if total_rows == 0:
            raise ValueError(f"No rows were found in {source_file}.")

        resolved_last_row = _resolve_last_row(total_rows=total_rows, configured_last_row=spec.last_row)
        header_anchor_row = max(spec.category_row or 0, spec.field_row)
        data_start_row = header_anchor_row + 1

        _validate_sheet_bounds(
            spec=spec,
            total_rows=total_rows,
            resolved_last_row=resolved_last_row,
            data_start_row=data_start_row,
            allow_header_only_empty=spec.last_row == -1,
        )

        header_rows = [
            list(row)
            for row in worksheet.iter_rows(min_row=1, max_row=header_anchor_row, values_only=True)
        ]
        if _streaming_header_may_need_merged_cell_expansion(raw_rows=header_rows, spec=spec):
            should_fallback_to_memory = True
        else:
            columns, duplicate_columns = _build_column_descriptors(raw_rows=header_rows, spec=spec)

            first_raw_row = next(
                worksheet.iter_rows(
                    min_row=data_start_row,
                    max_row=data_start_row,
                    max_col=_max_physical_column_index(columns),
                    values_only=True,
                ),
                None,
            )
            first_data_row = (
                _build_row_value_map(row_values=first_raw_row, columns=columns, source_file=source_file)
                if first_raw_row is not None
                else None
            )
    finally:
        workbook.close()

    if should_fallback_to_memory:
        return _build_in_memory_sheet_dataset(source_file=source_file, spec=spec, load_cache=load_cache)

    return SheetDataset(
        source_file=source_file,
        logical_sheet=spec.sheet,
        physical_sheet=physical_sheet,
        total_rows=total_rows,
        resolved_last_row=resolved_last_row,
        columns=columns,
        duplicate_columns=duplicate_columns,
        data_rows=[],
        data_row_numbers=[],
        first_data_row=first_data_row,
        first_data_row_number=data_start_row if first_data_row is not None else None,
        row_iterator_factory=lambda: _iter_workbook_sheet_rows(
            source_file=source_file,
            logical_sheet=spec.sheet,
            columns=columns,
            data_start_row=data_start_row,
            resolved_last_row=resolved_last_row,
            progress_callback=progress_callback,
        ),
    )


def _build_streaming_sheet_dataset_to_end(
    source_file: Path,
    spec: SourceSheetSpec,
    *,
    load_cache: SourceFileLoadCache | None = None,
    progress_callback: ProgressCallback | None = None,
) -> SheetDataset:
    """为 `last_row=-1` 的超大工作簿构建真正流式的数据集。"""

    from openpyxl import load_workbook

    should_fallback_to_memory = False
    workbook = load_workbook(source_file, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet, physical_sheet = _resolve_workbook_worksheet(workbook=workbook, logical_sheet=spec.sheet)
        _prepare_read_only_worksheet(worksheet, force_calculate=False)
        header_anchor_row = max(spec.category_row or 0, spec.field_row)
        data_start_row = header_anchor_row + 1

        if spec.field_row < 1:
            raise ValueError("field_row and last_row must resolve to positive row numbers.")
        if spec.category_row is not None and spec.category_row < 1:
            raise ValueError("category_row must resolve to a positive row number when provided.")

        header_rows = [
            list(row)
            for row in worksheet.iter_rows(min_row=1, max_row=header_anchor_row, values_only=True)
        ]
        if len(header_rows) < spec.field_row:
            raise ValueError("field_row exceeds the actual file row count.")
        if spec.category_row is not None and len(header_rows) < spec.category_row:
            raise ValueError("category_row exceeds the actual file row count.")

        if _streaming_header_may_need_merged_cell_expansion(raw_rows=header_rows, spec=spec):
            should_fallback_to_memory = True
        else:
            columns, duplicate_columns = _build_column_descriptors(raw_rows=header_rows, spec=spec)
            first_raw_row = next(
                worksheet.iter_rows(
                    min_row=data_start_row,
                    max_col=_max_physical_column_index(columns),
                    values_only=True,
                ),
                None,
            )
            if first_raw_row is None:
                first_data_row = None
                first_data_row_number = None
            else:
                first_data_row = _build_row_value_map(row_values=first_raw_row, columns=columns, source_file=source_file)
                first_data_row_number = data_start_row
    finally:
        workbook.close()

    if should_fallback_to_memory:
        return _build_in_memory_sheet_dataset(source_file=source_file, spec=spec, load_cache=load_cache)

    return SheetDataset(
        source_file=source_file,
        logical_sheet=spec.sheet,
        physical_sheet=physical_sheet,
        total_rows=data_start_row,
        resolved_last_row=data_start_row,
        columns=columns,
        duplicate_columns=duplicate_columns,
        data_rows=[],
        data_row_numbers=[],
        first_data_row=first_data_row,
        first_data_row_number=first_data_row_number,
        row_iterator_factory=lambda: _iter_workbook_sheet_rows_to_end(
            source_file=source_file,
            logical_sheet=spec.sheet,
            columns=columns,
            data_start_row=data_start_row,
            progress_callback=progress_callback,
        ),
    )


def find_rule_matches(dataset: SheetDataset, category: str | None, field: str) -> list[ColumnDescriptor]:
    """根据规则里的 `category + field` 在数据集里找列。"""

    if category is None:
        return [column for column in dataset.columns if column.field == field and not column.category]
    return [column for column in dataset.columns if column.field == field and column.category == category]


def split_rule_sheet_names(sheet_value: str) -> list[str]:
    """把规则中的 sheet 配置拆成可匹配的逻辑 sheet 名列表。"""

    return [sheet_name.strip() for sheet_name in sheet_value.split(",") if sheet_name.strip()]


def rule_applies_to_sheet(rule_sheet: str, logical_sheet: str) -> bool:
    """判断一条规则是否应当作用到当前逻辑 sheet。"""

    return logical_sheet in split_rule_sheet_names(rule_sheet)


def extract_rule_value(row_values: dict[tuple[str, str], Any], column: ColumnDescriptor) -> Any:
    """从单行数据中取出某个规则命中的值。"""

    return row_values.get((column.category or "", column.field))


def _load_raw_rows(
    source_file: Path,
    logical_sheet: str,
    *,
    load_cache: SourceFileLoadCache | None,
) -> tuple[list[list[Any]], str]:
    """读取物理文件并映射到逻辑 sheet。"""

    if load_cache is not None:
        return load_cache.get_raw_rows(source_file=source_file, logical_sheet=logical_sheet)

    suffix = source_file.suffix.lower()
    if suffix == ".csv":
        if logical_sheet != DEFAULT_LOGICAL_SHEET:
            raise ValueError("CSV files must use the DEFAULT logical sheet.")
        return _read_csv_rows(source_file), DEFAULT_LOGICAL_SHEET

    snapshot = _build_workbook_snapshot(source_file)
    return snapshot.resolve_rows(logical_sheet)


def _build_workbook_snapshot(source_file: Path) -> WorkbookSnapshot:
    """一次性加载工作簿，并缓存全部 sheet 的原始行数据。"""

    from openpyxl import load_workbook

    workbook = load_workbook(source_file, read_only=False, data_only=True, keep_links=False)
    try:
        raw_rows_by_sheet = {
            worksheet.title: _extract_raw_rows_from_worksheet(worksheet)
            for worksheet in workbook.worksheets
        }
        return WorkbookSnapshot(
            sheet_names=tuple(workbook.sheetnames),
            raw_rows_by_sheet=raw_rows_by_sheet,
        )
    finally:
        workbook.close()


def _read_csv_rows(source_file: Path) -> list[list[Any]]:
    """用约定编码集合读取 CSV，并过滤全空白行。"""

    decode_errors: list[str] = []
    for encoding in SUPPORTED_CSV_ENCODINGS:
        try:
            with source_file.open("r", encoding=encoding, newline="") as handle:
                return [
                    list(row)
                    for row in csv.reader(handle)
                    if any(cell.strip() for cell in row if isinstance(cell, str))
                ]
        except UnicodeDecodeError as exc:
            decode_errors.append(f"{encoding}: {exc}")

    supported = ", ".join(SUPPORTED_CSV_ENCODINGS)
    error_summary = "; ".join(decode_errors)
    raise ValueError(
        f"CSV file cannot be decoded with supported encodings ({supported}). {error_summary}"
    )


def _resolve_workbook_worksheet(workbook, logical_sheet: str):
    """把逻辑 sheet 映射到工作簿中的真实 worksheet。"""

    if logical_sheet == DEFAULT_LOGICAL_SHEET:
        if len(workbook.sheetnames) != 1:
            raise ValueError("DEFAULT logical sheet can only be used when the workbook has exactly one sheet.")
        physical_sheet = workbook.sheetnames[0]
    else:
        if logical_sheet not in workbook.sheetnames:
            raise KeyError(logical_sheet)
        physical_sheet = logical_sheet
    return workbook[physical_sheet], physical_sheet


def _count_worksheet_rows(worksheet) -> int:
    """在元数据缺失时兜底统计 worksheet 行数。"""

    return sum(1 for _ in worksheet.iter_rows(values_only=True))


def _extract_raw_rows_from_worksheet(worksheet) -> list[list[Any]]:
    """把 worksheet 展开成便于后续解析的二维数组。"""

    if not worksheet.merged_cells.ranges:
        return [list(row) for row in worksheet.iter_rows(values_only=True)]

    merged_value_map = _build_merged_value_map(worksheet)
    raw_rows: list[list[Any]] = []
    for row_index in range(1, worksheet.max_row + 1):
        row_values: list[Any] = []
        for column_index in range(1, worksheet.max_column + 1):
            row_values.append(
                merged_value_map.get((row_index, column_index), worksheet.cell(row=row_index, column=column_index).value)
            )
        raw_rows.append(row_values)
    return raw_rows


def _resolve_last_row(total_rows: int, configured_last_row: int) -> int:
    """将 `last_row` 解析成实际的结束行号。"""

    if configured_last_row == 0:
        raise ValueError("last_row cannot be zero.")
    if configured_last_row < 0:
        # 约定：-1 表示包含最后一行，-2 表示包含倒数第二行、不包含最后一行。
        return total_rows + configured_last_row + 1
    return configured_last_row


def _validate_sheet_bounds(
    *,
    spec: SourceSheetSpec,
    total_rows: int,
    resolved_last_row: int,
    data_start_row: int,
    allow_header_only_empty: bool = False,
) -> None:
    """统一校验 sheet 的行边界配置。"""

    if spec.field_row < 1 or resolved_last_row < 1:
        raise ValueError("field_row and last_row must resolve to positive row numbers.")
    if spec.category_row is not None and spec.category_row < 1:
        raise ValueError("category_row must resolve to a positive row number when provided.")
    if spec.field_row > total_rows:
        raise ValueError("field_row exceeds the actual file row count.")
    if spec.category_row is not None and spec.category_row > total_rows:
        raise ValueError("category_row exceeds the actual file row count.")
    if data_start_row > resolved_last_row:
        if allow_header_only_empty and data_start_row == total_rows + 1 and resolved_last_row == total_rows:
            return
        raise ValueError("The data range is empty after applying category_row, field_row, and last_row.")


def _normalize_cell_value(value: Any) -> str | None:
    """把单元格值归一成便于比较的文本。"""

    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _get_cell(rows: list[list[Any]], row_index: int, column_index: int) -> Any:
    """安全读取一个二维表格坐标。"""

    if row_index < 0 or row_index >= len(rows):
        return None
    row = rows[row_index]
    if column_index < 0 or column_index >= len(row):
        return None
    return row[column_index]


def _build_column_descriptors(
    *,
    raw_rows: list[list[Any]],
    spec: SourceSheetSpec,
) -> tuple[list[ColumnDescriptor], dict[tuple[str, str], list[int]]]:
    """从表头行中提取规则匹配所需的列描述。"""

    max_columns = max((len(row) for row in raw_rows), default=0)
    columns: list[ColumnDescriptor] = []
    seen_columns: dict[tuple[str, str], list[int]] = {}

    for column_index in range(max_columns):
        category, field = _resolve_column_header(raw_rows=raw_rows, spec=spec, column_index=column_index)
        if not field:
            continue

        key = (category or "", field)
        seen_columns.setdefault(key, []).append(column_index + 1)
        columns.append(ColumnDescriptor(column_index=column_index, category=category, field=field))

    _append_virtual_file_name_column(columns=columns, seen_columns=seen_columns)
    duplicate_columns = {key: indices for key, indices in seen_columns.items() if len(indices) > 1}
    return columns, duplicate_columns


def _build_row_value_map(
    *,
    row_values: tuple[Any, ...] | list[Any] | None,
    columns: list[ColumnDescriptor],
    source_file: Path | None = None,
) -> DataRow:
    """将单行二维表数据映射成规则引擎可消费的字典。"""

    normalized_row = list(row_values) if row_values is not None else []
    mapped_row: DataRow = {}
    for column in columns:
        if column.column_index < 0 and column.field == FILE_NAME_VIRTUAL_FIELD:
            mapped_row[(column.category or "", column.field)] = source_file.stem if source_file is not None else None
            continue
        raw_value = normalized_row[column.column_index] if column.column_index < len(normalized_row) else None
        mapped_row[(column.category or "", column.field)] = _normalize_source_cell_value(raw_value)
    return mapped_row


def _is_effectively_blank_data_row(row_values: DataRow) -> bool:
    """判断一行数据在忽略虚拟列后是否全为空白。"""

    for (category, field), value in row_values.items():
        if category == "" and field == FILE_NAME_VIRTUAL_FIELD:
            continue
        if value not in (None, ""):
            return False
    return True


def _iter_workbook_sheet_rows(
    *,
    source_file: Path,
    logical_sheet: str,
    columns: list[ColumnDescriptor],
    data_start_row: int,
    resolved_last_row: int,
    progress_callback: ProgressCallback | None = None,
) -> Iterator[tuple[int, DataRow]]:
    """流式遍历单个工作簿 sheet 的数据行。"""

    from openpyxl import load_workbook

    workbook = load_workbook(source_file, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet, _ = _resolve_workbook_worksheet(workbook=workbook, logical_sheet=logical_sheet)
        _prepare_read_only_worksheet(worksheet, force_calculate=True)
        for row_number, row_values in enumerate(
            worksheet.iter_rows(
                min_row=data_start_row,
                max_row=resolved_last_row,
                max_col=_max_physical_column_index(columns),
                values_only=True,
            ),
            start=data_start_row,
        ):
            mapped_row = _build_row_value_map(row_values=row_values, columns=columns, source_file=source_file)
            if _is_effectively_blank_data_row(mapped_row):
                continue
            if progress_callback is not None:
                progress_callback(row_number)
            yield row_number, mapped_row
    finally:
        workbook.close()


def _iter_workbook_sheet_rows_to_end(
    *,
    source_file: Path,
    logical_sheet: str,
    columns: list[ColumnDescriptor],
    data_start_row: int,
    progress_callback: ProgressCallback | None = None,
) -> Iterator[tuple[int, DataRow]]:
    """流式遍历 worksheet 到物理文件末尾。"""

    from openpyxl import load_workbook

    workbook = load_workbook(source_file, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet, _ = _resolve_workbook_worksheet(workbook=workbook, logical_sheet=logical_sheet)
        _prepare_read_only_worksheet(worksheet, force_calculate=False)
        for row_number, row_values in enumerate(
            worksheet.iter_rows(
                min_row=data_start_row,
                max_col=_max_physical_column_index(columns),
                values_only=True,
            ),
            start=data_start_row,
        ):
            mapped_row = _build_row_value_map(row_values=row_values, columns=columns, source_file=source_file)
            if _is_effectively_blank_data_row(mapped_row):
                continue
            if progress_callback is not None:
                progress_callback(row_number)
            yield row_number, mapped_row
    finally:
        workbook.close()


def _prepare_read_only_worksheet(worksheet, *, force_calculate: bool) -> None:
    """修正只读模式下可能失真的 worksheet 维度元数据。"""

    if hasattr(worksheet, "reset_dimensions"):
        worksheet.reset_dimensions()
        if force_calculate and hasattr(worksheet, "calculate_dimension"):
            worksheet.calculate_dimension(force=True)


def _max_physical_column_index(columns: list[ColumnDescriptor]) -> int | None:
    """Return the rightmost physical source column required by rules.

    Args:
        columns: Resolved source columns, including virtual columns such as
            ``FILE_NAME`` whose ``column_index`` is negative.

    Returns:
        1-based maximum column index for ``openpyxl.iter_rows(max_col=...)``.
        Returns None when there are no physical columns.
    """

    physical_indices = [column.column_index for column in columns if column.column_index >= 0]
    if not physical_indices:
        return None
    return max(physical_indices) + 1


def _streaming_header_may_need_merged_cell_expansion(
    *,
    raw_rows: list[list[Any]],
    spec: SourceSheetSpec,
) -> bool:
    """Return whether streaming header parsing may lose merged-cell categories.

    Args:
        raw_rows: Header rows read from a read-only worksheet.
        spec: Source sheet structure from the rule bundle.

    Returns:
        True when a two-row header has a field value but no category value in
        the configured category row. In read-only mode openpyxl does not expose
        merged-cell ranges, so this shape may mean the category value should
        have been expanded from a merged cell by the in-memory parser.
    """

    if spec.category_row is None:
        return False
    category_row_index = spec.category_row - 1
    field_row_index = spec.field_row - 1
    max_columns = max((len(row) for row in raw_rows), default=0)
    for column_index in range(max_columns):
        field_value = _normalize_cell_value(_get_cell(raw_rows, field_row_index, column_index))
        category_value = _normalize_cell_value(_get_cell(raw_rows, category_row_index, column_index))
        if field_value and category_value is None:
            return True
    return False


def _resolve_column_header(
    raw_rows: list[list[Any]],
    spec: SourceSheetSpec,
    column_index: int,
) -> tuple[str | None, str | None]:
    """根据 source spec 解析单列的 category / field。"""

    current_row_value = _normalize_cell_value(_get_cell(raw_rows, spec.field_row - 1, column_index))
    if spec.category_row is None:
        return None, current_row_value

    previous_row_index = spec.field_row - 2
    previous_row_value = _normalize_cell_value(_get_cell(raw_rows, previous_row_index, column_index))

    # 跨两行合并的“无分类直属字段”：上一行和当前行是同一个值。
    if current_row_value and previous_row_value and current_row_value == previous_row_value:
        return None, current_row_value

    # 常规双层表头：上一行为分类，当前行为字段。
    if current_row_value:
        return previous_row_value, current_row_value

    # 兜底：当前行为空但上一行有值时，仍视为无分类字段。
    return None, previous_row_value


def _build_merged_value_map(worksheet) -> dict[tuple[int, int], Any]:
    """把合并单元格范围展开成逐坐标值映射。"""

    merged_value_map: dict[tuple[int, int], Any] = {}
    for merged_range in worksheet.merged_cells.ranges:
        top_left_value = worksheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        for row_index in range(merged_range.min_row, merged_range.max_row + 1):
            for column_index in range(merged_range.min_col, merged_range.max_col + 1):
                merged_value_map[(row_index, column_index)] = top_left_value
    return merged_value_map


def _append_virtual_file_name_column(
    *,
    columns: list[ColumnDescriptor],
    seen_columns: dict[tuple[str, str], list[int]],
) -> None:
    """追加内建的来源文件虚拟列。"""

    key = ("", FILE_NAME_VIRTUAL_FIELD)
    if key in seen_columns:
        return
    seen_columns[key] = [-1]
    columns.append(ColumnDescriptor(column_index=-1, category=None, field=FILE_NAME_VIRTUAL_FIELD))


def _normalize_source_cell_value(value: Any) -> Any:
    """清理常见的源文件文本噪音。"""

    if isinstance(value, str):
        return value.lstrip("`")
    return value
